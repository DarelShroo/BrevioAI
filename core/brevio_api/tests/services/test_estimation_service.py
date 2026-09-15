from pathlib import Path
from tempfile import NamedTemporaryFile
from types import SimpleNamespace
from typing import Any, List
from unittest.mock import AsyncMock, MagicMock, patch

import pytest
from bson import ObjectId
from docx import Document
from fastapi.testclient import TestClient
from pydantic import ValidationError

from core.brevio.enums.category import CategoryType
from core.brevio.enums.language import LanguageType
from core.brevio.enums.output_format_type import OutputFormatType
from core.brevio.enums.summary_level import SummaryLevel
from core.brevio_api.__main__ import app
from core.brevio_api.dependencies.api_key_dependency import verify_api_key
from core.brevio_api.dependencies.estimation_service_dependency import (
    get_estimation_service,
)
from core.brevio_api.dependencies.user_dependency import get_current_user
from core.brevio_api.models.brevio.document_processing_request import (
    DocumentProcessingRequest,
)
from core.brevio_api.models.brevio.responses.estimate_documents_response import (
    EstimationWarning,
)
from core.brevio_api.services.estimation.document_extraction_service import (
    DocumentExtractionService,
    ExtractedDocument,
)
from core.brevio_api.services.estimation.estimation_service import EstimationService
from core.brevio_api.services.estimation.tokenizer_service import (
    TokenCountResult,
    TokenizerService,
)
from core.shared.enums.model import ModelType
from core.shared.utils.context_validator import ContextValidator
from core.shared.utils.model_capabilities import get_model_capabilities
from core.shared.utils.output_token_policy import OutputTokenPolicy
from core.shared.utils.pricing_registry import (
    calculate_credits,
    calculate_provider_cost,
    convert_to_platform_currency,
    get_pricing_snapshot,
)

client = TestClient(app)


@pytest.fixture(autouse=True)
def clear_overrides() -> None:
    yield
    app.dependency_overrides.clear()


@pytest.fixture
def document_request() -> DocumentProcessingRequest:
    return DocumentProcessingRequest(
        language=LanguageType.SPANISH,
        model=ModelType.GPT_4O_MINI,
        category=CategoryType.EDUCATION.value,
        style="guide",
        format=OutputFormatType.MARKDOWN,
        summary_level=SummaryLevel.CONCISE,
    )


@pytest.mark.asyncio
async def test_tokenizer_exact_and_approx(monkeypatch: pytest.MonkeyPatch) -> None:
    service = TokenizerService()
    exact = service.count_tokens("hello world", ModelType.GPT_4O_MINI)
    assert exact.method == "exact"
    assert exact.tokens > 0

    def raising_get_encoder(model: ModelType) -> Any:
        raise RuntimeError("missing tokenizer")

    monkeypatch.setattr(
        "core.brevio_api.services.estimation.tokenizer_service.get_encoder",
        raising_get_encoder,
    )
    approx = service.count_tokens("x" * 100, ModelType.GPT_4O_MINI)
    assert approx.method == "approximation"
    assert approx.confidence == "medium"
    assert approx.tokens > 0


@pytest.mark.asyncio
async def test_document_extraction_txt_and_markdown() -> None:
    service = DocumentExtractionService()

    txt = await service._extract_single_document("notes.txt", b"Line 1\n\nLine 2")
    md = await service._extract_single_document("readme.md", b"# Title\n\nBody")

    assert txt.extraction_status == "complete"
    assert md.extraction_status == "complete"
    assert "Line 1" in txt.text
    assert "Title" in md.text


@pytest.mark.asyncio
async def test_document_extraction_docx() -> None:
    service = DocumentExtractionService()
    with NamedTemporaryFile(suffix=".docx", delete=False) as tmp_file:
        tmp_path = Path(tmp_file.name)
    try:
        doc = Document()
        doc.add_paragraph("Hello Brevio")
        doc.save(tmp_path)
        result = await service._extract_docx("example.docx", tmp_path.read_bytes())
        assert result.extraction_status == "complete"
        assert "Hello Brevio" in result.text
    finally:
        tmp_path.unlink(missing_ok=True)


@pytest.mark.asyncio
async def test_document_extraction_pdf_statuses() -> None:
    service = DocumentExtractionService()
    blank_page = MagicMock()
    blank_page.extract_text.return_value = ""
    blank_page.get.return_value = None

    text_page = MagicMock()
    text_page.extract_text.return_value = "Extracted text"
    text_page.get.return_value = None

    reader = MagicMock()
    reader.pages = [text_page, blank_page]

    with patch.object(service, "_load_pdf_reader", return_value=reader):
        partial = await service._extract_pdf("mixed.pdf", b"%PDF-1.4")

    assert partial.extraction_status == "partial"
    assert partial.pages_total == 2
    assert partial.pages_with_text == 1
    assert partial.pages_without_text == 1

    reader_empty = MagicMock()
    reader_empty.pages = [blank_page]
    with patch.object(
        service, "_load_pdf_reader", return_value=reader_empty
    ), patch.object(
        service,
        "_page_looks_scanned",
        return_value=False,
    ):
        empty = await service._extract_pdf("empty.pdf", b"%PDF-1.4")
    assert empty.extraction_status == "empty"

    reader_scanned = MagicMock()
    reader_scanned.pages = [blank_page]
    with patch.object(
        service, "_load_pdf_reader", return_value=reader_scanned
    ), patch.object(
        service,
        "_page_looks_scanned",
        return_value=True,
    ):
        scanned = await service._extract_pdf("scanned.pdf", b"%PDF-1.4")
    assert scanned.extraction_status == "ocr_required"


def test_context_validator_and_output_policy() -> None:
    policy = OutputTokenPolicy()
    from core.brevio.enums.language import LanguageType

    budget = policy.budget(
        ModelType.GPT_4O_MINI,
        SummaryLevel.DETAILED,
        OutputFormatType.MARKDOWN,
        "guide",
        language=LanguageType.ENGLISH,
    )
    capabilities = get_model_capabilities(ModelType.GPT_4O_MINI)
    result = ContextValidator().validate(1000, budget, capabilities)
    assert result.can_process is True
    assert result.reason is None
    assert budget.minimum <= budget.expected <= budget.maximum


def test_pricing_and_credits() -> None:
    snapshot = get_pricing_snapshot(ModelType.GPT_4O_MINI)
    expected_provider_cost, max_provider_cost = calculate_provider_cost(
        10_000,
        1_000,
        2_000,
        snapshot,
    )
    platform_expected = convert_to_platform_currency(expected_provider_cost, snapshot)
    platform_max = convert_to_platform_currency(max_provider_cost, snapshot)
    assert platform_expected <= platform_max
    assert calculate_credits(platform_expected, snapshot) >= snapshot.minimum_credits
    assert calculate_credits(platform_max, snapshot) >= calculate_credits(
        platform_expected, snapshot
    )

    local_snapshot = get_pricing_snapshot(ModelType.QWEN3_8B)
    assert local_snapshot.input_price_per_million == 0.0
    assert calculate_credits(0.0, local_snapshot) == 0
    assert calculate_credits(1.0, local_snapshot) == 0


@pytest.mark.asyncio
async def test_estimation_service_aggregates_per_document(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    service = EstimationService()
    request = DocumentProcessingRequest(
        language=LanguageType.SPANISH,
        model=ModelType.GPT_4O_MINI,
        category=CategoryType.EDUCATION.value,
        style="guide",
        format=OutputFormatType.MARKDOWN,
        summary_level=SummaryLevel.CONCISE,
    )

    extracted_documents = [
        ExtractedDocument(
            filename="guide.pdf",
            extension="pdf",
            text="alpha beta gamma",
            extraction_status="complete",
            pages_total=1,
            pages_with_text=1,
            pages_without_text=0,
        ),
        ExtractedDocument(
            filename="notes.md",
            extension="md",
            text="delta epsilon",
            extraction_status="complete",
        ),
    ]

    async def fake_build_prompt(*args: Any, **kwargs: Any) -> str:
        return "prompt"

    def fake_count_tokens(text: str, model: ModelType) -> TokenCountResult:
        mapping = {
            "prompt": TokenCountResult(11, "exact", "high", 0.0),
            service._prompt_estimation.structural_wrapper_text(""): TokenCountResult(
                4, "exact", "high", 0.0
            ),
            "alpha beta gamma": TokenCountResult(3, "exact", "high", 0.0),
            "delta epsilon": TokenCountResult(2, "exact", "high", 0.0),
        }
        return mapping[text]

    monkeypatch.setattr(service._prompt_estimation, "build_prompt", fake_build_prompt)
    monkeypatch.setattr(
        service._extraction_service,
        "extract_documents",
        AsyncMock(return_value=extracted_documents),
    )
    monkeypatch.setattr(service._tokenizer, "count_tokens", fake_count_tokens)

    outcome = await service.estimate(
        request, [(doc.filename, b"content") for doc in extracted_documents]
    )

    assert outcome.can_process is True
    assert outcome.total_document_tokens == 5
    assert outcome.prompt_template_tokens == 22
    assert outcome.structural_tokens == 8
    assert outcome.tokenization_method == "exact"
    assert outcome.estimation_id == outcome.estimation_id


def test_estimate_documents_endpoint(monkeypatch: pytest.MonkeyPatch) -> None:
    mock_service = MagicMock()
    mock_service.estimate = AsyncMock(
        return_value=SimpleNamespace(
            can_process=True,
            reason=None,
            prompt_template_tokens=10,
            total_document_tokens=20,
            structural_tokens=4,
            total_input_tokens=34,
            output_minimum_tokens=100,
            output_expected_tokens=150,
            output_max_tokens=200,
            document_token_counts=[
                (
                    ExtractedDocument(
                        filename="guide.pdf",
                        extension="pdf",
                        text="alpha",
                        extraction_status="complete",
                        pages_total=1,
                        pages_with_text=1,
                        pages_without_text=0,
                    ),
                    20,
                )
            ],
            tokenization_method="exact",
            tokenization_confidence="high",
            tokenization_error_margin=0.0,
            provider_expected_cost=0.001,
            provider_max_cost=0.002,
            platform_expected_cost=0.003,
            platform_max_cost=0.004,
            expected_credits=1,
            maximum_credits=1,
            chunking_required=False,
            chunking_calls=1,
            estimation_id="abc123",
            estimated_at="2026-09-09T00:00:00Z",
            warnings=[],
            request=DocumentProcessingRequest(
                language=LanguageType.SPANISH,
                model=ModelType.GPT_4O_MINI,
                category=CategoryType.EDUCATION.value,
                style="guide",
                format=OutputFormatType.MARKDOWN,
                summary_level=SummaryLevel.CONCISE,
            ),
        )
    )

    app.dependency_overrides[verify_api_key] = lambda: "test-key"
    app.dependency_overrides[get_current_user] = lambda: str(ObjectId())
    app.dependency_overrides[get_estimation_service] = lambda: mock_service

    files = {"files": ("guide.pdf", b"fake-pdf", "application/pdf")}
    data = {
        "language": "SPANISH",
        "model": "gpt-4o-mini",
        "category": "education",
        "style": "guide",
        "format": "markdown",
        "summary_level": "concise",
    }

    response = client.post("/brevio/estimate-documents", files=files, data=data)
    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] == "success"
    assert payload["data"]["can_process"] is True
    assert payload["data"]["pricing"]["expected_cost"] == 0.003
