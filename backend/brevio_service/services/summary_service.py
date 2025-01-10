import os
import time
from openai import OpenAI
from ..constants.constants import Constants
from ..constants.summary_messages import SummaryMessages
from ..enums.role import RoleType
from ..managers.directory_manager import DirectoryManager
from ..models.config_model import ConfigModel as Config
from ..models.response_model import SummaryResponse
import markdown
from docx import Document
from bs4 import BeautifulSoup

class Summary:
    def __init__(self, transcription_path, summary_path, content, config: Config):
        self.transcription_path = transcription_path
        self.summary_path = summary_path
        self.content = content
        self.config = config

        self.api_key = self.config.api_key
        self.max_tokens = self.config.max_tokens
        self.tokens_per_minute = self.config.tokens_per_minute
        self.temperature = self.config.temperature

        self.client = OpenAI(api_key=self.api_key)
        self.model = self.config.model.value

    def generate_summary_chunk(self, total_tokens_used, chunk):
        if total_tokens_used + int(self.max_tokens) > int(self.tokens_per_minute):
            print(SummaryMessages.WAITING)
            time.sleep(60)
            total_tokens_used = 0

        response = self.client.chat.completions.create(
            model=self.model,
            messages=[
                {"role": RoleType.SYSTEM.value, "content": self.content},
                {"role": RoleType.USER.value, "content": chunk}
            ],
            max_tokens=int(self.max_tokens),
            temperature=float(self.temperature)
        )

        summary = response.choices[0].message.content.strip()
        total_tokens_used += int(self.max_tokens)
        print(summary)
        return summary, total_tokens_used

    def markdown_to_docx(self, md_file_path, docx_file_path):
        try:
            with open(md_file_path, 'r', encoding='utf-8') as md_file:
                md_content = md_file.read()

            html_content = markdown.markdown(md_content)

            doc = Document()

            def add_html_to_docx(html):
                soup = BeautifulSoup(html, 'html.parser')
                for element in soup:
                    if element.name == 'h1':
                        doc.add_heading(element.text, level=1)
                    elif element.name == 'h2':
                        doc.add_heading(element.text, level=2)
                    elif element.name == 'h3':
                        doc.add_heading(element.text, level=3)
                    elif element.name == 'p':
                        doc.add_paragraph(element.text)
                    elif element.name == 'ul':
                        for li in element.find_all('li'):
                            doc.add_paragraph(li.text, style='List Bullet')

            add_html_to_docx(html_content)

            doc.save(docx_file_path)
            print(f"Archivo Word guardado en {docx_file_path}")

        except Exception as e:
            print(f"Error al convertir Markdown a Word: {e}")
            raise

    def generate_summary(self):
        try:
            directory_manager = DirectoryManager(self.config)
            directory_manager.validate_paths(self.transcription_path)
            
            transcription = directory_manager.read_transcription(self.transcription_path)

            chunks = [
                transcription[i:i + int(self.max_tokens)] 
                for i in range(0, len(transcription), int(self.max_tokens))
            ]

            summary = ""
            total_tokens_used = 0
            for chunk in chunks:
                chunk_summary, total_tokens_used = self.generate_summary_chunk(total_tokens_used, chunk)
                summary += chunk_summary

            directory_manager.write_summary(summary, self.summary_path)

            docx_path = self.summary_path.replace('.md', '.docx')
            self.markdown_to_docx(self.summary_path, docx_path)

            return SummaryResponse(
                success=True, 
                summary=str(summary),
                message=str(SummaryMessages.WRITE_SUMMARY.format(self.summary_path))
            )

        except Exception as e:
            return SummaryResponse(
                success=False, 
                summary=f"Error occurred: {str(e)}",
                message=SummaryMessages.ERROR_GENERATING_SUMMARY.format(str(e))
            )
